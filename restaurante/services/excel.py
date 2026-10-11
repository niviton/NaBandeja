"""Exportação do relatório de vendas para Excel (.xlsx)."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill('solid', fgColor='1E293B')
HEADER_FONT = Font(bold=True, color='FFFFFF')
MONEY = 'R$ #,##0.00'


def _sheet(workbook, title, headers, rows, money_columns=(), percent_columns=()):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal='center')
    for row in rows:
        sheet.append(list(row))
    for column in money_columns:
        for cell in sheet[get_column_letter(column)][1:]:
            cell.number_format = MONEY
    for column in percent_columns:
        for cell in sheet[get_column_letter(column)][1:]:
            cell.number_format = '0.0%'
    for index, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(row[index - 1] or '')) for row in rows] or [10])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 10), 60)
    sheet.freeze_panes = 'A2'
    return sheet


def sales_workbook(store_name, period_label, summary, history_orders):
    workbook = Workbook()
    resume = workbook.active
    resume.title = 'Resumo'
    resume.append([f'{store_name} - Relatório de vendas'])
    resume['A1'].font = Font(bold=True, size=14)
    resume.append(['Período', period_label])
    for label, value in [
        ('Faturamento', summary['revenue']),
        ('Pedidos', summary['orders']),
        ('Ticket médio', summary['avg_ticket']),
        ('Descontos', summary['discounts']),
        ('Taxas de entrega', summary['delivery_fees']),
        ('Custo dos produtos', summary['items_cost']),
        ('Lucro bruto', summary['gross_profit']),
    ]:
        resume.append([label, value])
        if label != 'Pedidos':
            resume.cell(resume.max_row, 2).number_format = MONEY
    resume.column_dimensions['A'].width = 22
    resume.column_dimensions['B'].width = 20

    _sheet(
        workbook, 'Vendas',
        ['Pedido', 'Fechado em', 'Cliente/Mesa', 'Tipo', 'Garçom', 'Status', 'Pagamento', 'Itens', 'Desconto', 'Entrega', 'Total'],
        [
            (
                order.id,
                order.closed_at.astimezone().strftime('%d/%m/%Y %H:%M') if order.closed_at else '',
                order.display_name,
                order.get_order_type_display(),
                order.waiter.username if order.waiter else '',
                order.get_status_display(),
                order.get_payment_method_display() if order.payment_method else '',
                ' | '.join(f'{item.quantity}x {item.name}' for item in order.items.all()),
                order.discount,
                order.delivery_fee,
                order.amount_due if order.status == 'finalizado' else 0,
            )
            for order in history_orders
        ],
        money_columns=(9, 10, 11),
    )
    _sheet(
        workbook, 'Produtos',
        ['Produto', 'Quantidade', 'Faturamento', 'Custo', 'Lucro', 'Margem'],
        [
            (row['product_name'], row['total_qty'], row['total_revenue'], row['total_cost'] if row['has_cost'] else None,
             row['profit'] if row['has_cost'] else None, float(row['margin']) / 100 if row['margin'] is not None else None)
            for row in summary['by_product']
        ],
        money_columns=(3, 4, 5), percent_columns=(6,),
    )
    _sheet(
        workbook, 'Pagamentos', ['Forma', 'Quantidade', 'Total'],
        [(row['method'], row['count'], row['total']) for row in summary['by_payment']],
        money_columns=(3,),
    )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
